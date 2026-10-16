import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import csv
import datetime
import os
import shutil

# --- 1. KEYWORD DICTIONARY FROM CSV ---
csv_path = 'keyword research CLEANWORX v2.csv'
keywords_dict = {}
keywords_list = []

with open(csv_path, encoding='utf-8-sig', errors='replace') as f:
    r = csv.reader(f)
    headers = next(r)
    for row in r:
        if row and len(row) >= 7:
            no = int(row[0]) if row[0].isdigit() else row[0]
            kw = row[1].strip()
            vol = float(row[2]) if row[2].replace('.', '', 1).isdigit() else 0.0
            cpc = float(row[3]) if row[3].replace('.', '', 1).isdigit() else 0.0
            pd = float(row[4]) if row[4].replace('.', '', 1).isdigit() else 0.0
            sd = float(row[5]) if row[5].replace('.', '', 1).isdigit() else 0.0
            cluster = row[6].strip()
            keywords_dict[kw] = vol
            keywords_list.append((no, kw, vol, cpc, pd, sd, cluster))

print(f"Loaded {len(keywords_list)} keywords.")

# --- 2. 52-WEEK CALENDAR DATA ---
# Each tuple: (
#   week_num (1-52),
#   date_str (YYYY-MM-DD),
#   objetivo (Spanish),
#   keyword (must match CSV exactly),
#   titulo_gbp (English),
#   texto_gbp (English, ready to publish),
#   foto_video (Spanish),
#   boton (GBP button: 'Learn more', 'Book', 'Get directions'),
#   destino (Canonical URL with CleanWorx UTM structure or Maps link),
#   revision (Spanish QC check),
#   categoria_resumen (for summary categorization)
# )

posts_data = [
    # Week 1: 2026-10-19
    (
        1, "2026-10-19",
        "Posicionar la protección cerámica antes del invierno y las heladas de NJ",
        "ceramic coating",
        "Prepare Your Paint for Winter with Ceramic Coating",
        "Winter in New Jersey brings corrosive road salt, brine slush, and freezing moisture that eat into unprotected clear coats. A professional System X ceramic coating bonds directly with your vehicle's paint, creating an ultra-slick hydrophobic barrier that prevents harsh road grime from bonding. Protect your vehicle at our 19 E. Henry Street studio in Basking Ridge before the first freeze.",
        "Foto real en primer plano de un vehículo recién tratado con gotas de agua perfectamente esféricas sobre el capó en el taller de CleanWorx.",
        "Learn more",
        "https://www.cleanworxnj.com/ceramic-coating?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_ceramic_coating_winter_prep_20261019",
        "Confirmar disponibilidad de turnos de coating en el taller de 19 E. Henry St.",
        "Protección cerámica (Ceramic Coating)"
    ),
    # Week 2: 2026-10-26
    (
        2, "2026-10-26",
        "Captar conductores que necesitan limpieza profunda interior antes del frío",
        "car interior cleaning",
        "Deep Interior Cleaning & Stain Extraction in Basking Ridge",
        "Before cold winter weather keeps your vehicle's windows closed for months, refresh your cabin with professional interior detailing. CleanWorx provides deep carpet extraction, upholstery shampooing, steam cleaning for vents and consoles, and conditioning for high-touch leather. Drop off your car at 19 E. Henry Street or request our self-contained mobile detailing unit.",
        "Foto dividida 50/50 de alfombra manchada frente a la sección limpia tras pasar el extractor de inyección y succión.",
        "Book",
        "https://www.cleanworxnj.com/interior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_interior_deep_cleaning_20261026",
        "Verificar que los precios de catálogo Square para interior coincidan con la cotización.",
        "Detallado interior (Interior Detailing)"
    ),
    # Week 3: 2026-11-02
    (
        3, "2026-11-02",
        "Explicar beneficios del polarizado cerámico para el resplandor solar otoñal e invernal",
        "window tint installation",
        "Precision Window Tint Installation in Basking Ridge, NJ",
        "Low autumn and winter sun angles create dangerous glare across New Jersey highways. CleanWorx offers precision-cut Carbon and Ceramic window film installed inside our clean, dust-controlled Basking Ridge studio. Enjoy reduced eye strain, 99% UV rejection, and enhanced cabin comfort backed by manufacturer warranties up to 10 years. Inquire today.",
        "Foto de un sedán premium recién polarizado en el área de trabajo con luz cenital destacando la nitidez del film.",
        "Learn more",
        "https://www.cleanworxnj.com/window-tinting?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_window_tint_installation_20261102",
        "Recordar las garantías oficiales: 2 años en Carbon film y 10 años en Ceramic film.",
        "Polarizado automotriz (Window Tinting)"
    ),
    # Week 4: 2026-11-09
    (
        4, "2026-11-09",
        "Mostrar la diferencia entre encerar y corregir la pintura con pulido profesional",
        "paint correction",
        "Remove Swirl Marks with Precision Machine Paint Correction",
        "Waxes only temporarily fill paint imperfections; true paint correction permanently levels the clear coat through controlled multi-stage polishing. At CleanWorx, we measure paint depth before eliminating swirl marks, wash scratches, and oxidation to restore mirror-like clarity. Standalone services start from $350 at our 19 E. Henry Street shop in Basking Ridge.",
        "Foto o video 50/50 iluminado con luz de inspección Scangrip mostrando la mitad pulida libre de remolinos frente a la mitad rayada.",
        "Book",
        "https://www.cleanworxnj.com/paint-correction?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_paint_correction_swirl_removal_20261109",
        "No prometer eliminación total de rayones que traspasen la capa transparente (clear coat).",
        "Corrección de pintura (Paint Correction)"
    ),
    # Week 5: 2026-11-16
    (
        5, "2026-11-16",
        "Promover la conveniencia del servicio móvil a domicilio o lugar de trabajo",
        "mobile auto detailing",
        "Self-Contained Mobile Auto Detailing Brought to Your Door",
        "Busy week in Somerset or Morris County? CleanWorx brings high-end detailing directly to your home driveway or workplace. Our custom mobile rig carries onboard spot-free deionized water and independent power—requiring zero hookups from your property. Available for detailing and tint appointments throughout Basking Ridge, Bernardsville, Bedminster, and beyond.",
        "Foto de la camioneta móvil de CleanWorx operando en la entrada de una residencia, con el equipo organizado.",
        "Learn more",
        "https://www.cleanworxnj.com/mobile-auto-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_mobile_auto_detailing_van_20261116",
        "Confirmar la regla de tarifa móvil: $50 en servicios menores a $400; sin costo adicional si es $400 o más.",
        "Servicio móvil a domicilio (Mobile Detailing)"
    ),
    # Week 6: 2026-11-23
    (
        6, "2026-11-23",
        "Enfatizar la descontaminación exterior antes del despliegue de salmuera invernal",
        "exterior car detailing",
        "Exterior Hand Detailing & Paint Decontamination",
        "Abrasive automated brush car washes scratch your finish and leave road films behind. CleanWorx exterior detailing includes a gentle multi-bucket hand wash, chemical iron decontamination, clay bar treatment to lift embedded road grit, and long-lasting paint sealant. Keep your exterior protected against winter road brine in Basking Ridge.",
        "Foto del proceso de descontaminación química mostrando la reacción púrpura (iron remover) en pintura y rines.",
        "Book",
        "https://www.cleanworxnj.com/exterior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_exterior_detailing_prep_20261123",
        "No usar la palabra 'car wash' como término de servicio principal; enfatizar detallado profesional a mano.",
        "Detallado exterior (Exterior Detailing)"
    ),
    # Week 7: 2026-11-30
    (
        7, "2026-11-30",
        "Reforzar la ubicación local y la experiencia del equipo en 19 E. Henry St",
        "car detailing shop near me",
        "Dedicated Detailing Studio in Basking Ridge, NJ",
        "Looking for an independent detailing shop that treats your car like their own? CleanWorx operates from our dedicated facility at 19 E. Henry Street in downtown Basking Ridge. Founded in 2019, our team—Vito, Melqui, and Hemza—personally manages every service with specialized equipment and climate-controlled bays. Stop by or book your appointment online.",
        "Foto del frente del taller en 19 E. Henry St con el equipo de CleanWorx junto a vehículos de clientes.",
        "Get directions",
        "https://maps.google.com/?q=CleanWorx+Auto+Detailing+19+E+Henry+Street+Basking+Ridge+NJ+07920",
        "Verificar que la dirección y el pin de Google Maps sean exactamente 19 E. Henry St, Basking Ridge, NJ 07920.",
        "Taller, equipo y reputación (Basking Ridge Studio & Trust)"
    ),
    # Week 8: 2026-12-07
    (
        8, "2026-12-07",
        "Promover restauración de faros por seguridad ante los días más oscuros del año",
        "headlight restoration",
        "Clear Headlight Restoration for Safer Winter Night Driving",
        "Yellowed, hazy headlights can reduce nighttime visibility by more than 75% on dark New Jersey roadways. CleanWorx offers professional multi-stage wet sanding, compounding, and protective UV sealing to restore crystal clarity to oxidized lenses. Available as a $75 add-on to any detail or as a $125 standalone service at our Basking Ridge shop.",
        "Foto comparativa de faro antes (amarillento y opaco) y después (completamente transparente como nuevo).",
        "Learn more",
        "https://www.cleanworxnj.com/add-ons?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_headlight_restoration_winter_20261207",
        "Mantener los precios aprobados de catálogo: $75 como add-on, $125 como servicio independiente.",
        "Servicios suplementarios (Odor, Engine, Headlights)"
    ),
    # Week 9: 2026-12-14
    (
        9, "2026-12-14",
        "Explicar las opciones de durabilidad de System X Ceramic Coating (1, 3 y 6 años)",
        "car ceramic coating",
        "System X Certified Ceramic Coating Packages",
        "Not all paint coatings offer the same durability. As a System X certified installer, CleanWorx offers 1-year, 3-year, and 6-year ceramic coating packages engineered to resist salt spray, road oils, and UV breakdown. Every installation includes rigorous prep and single-stage paint enhancement to ensure maximum bonding and gloss. Schedule your consultation today.",
        "Foto del instalador aplicando el coating con aplicador de microfibra sobre panel preparado en taller iluminado.",
        "Learn more",
        "https://www.cleanworxnj.com/ceramic-coating?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_system_x_ceramic_packages_20261214",
        "Confirmar la certificación System X y recordar no afirmar garantía de fábrica hasta validar condiciones del cliente.",
        "Protección cerámica (Ceramic Coating)"
    ),
    # Week 10: 2026-12-21
    (
        10, "2026-12-21",
        "Captar limpiezas de tapicería y asientos para los viajes navideños y visitas familiares",
        "car carpet cleaning",
        "Carpet & Upholstery Deep Cleaning for Holiday Road Trips",
        "Hosting family or traveling for the holidays? Make sure your vehicle's interior is clean, fresh, and welcoming. Our heated extractor penetrates deep into automotive carpet fibers to flush out ground-in dirt, winter moisture, and beverage stains. Drop off your vehicle at 19 E. Henry Street in Basking Ridge for our full interior service.",
        "Foto del tanque de recuperación del extractor mostrando el agua oscura extraída frente a las alfombras limpias.",
        "Book",
        "https://www.cleanworxnj.com/interior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_holiday_carpet_cleaning_20261221",
        "Confirmar horarios de atención durante la semana festiva (recordar cerrado los domingos).",
        "Detallado interior (Interior Detailing)"
    ),
    # Week 11: 2026-12-28
    (
        11, "2026-12-28",
        "Ofrecer remoción de polarizado viejo, morado o con burbujas",
        "windows tint removal",
        "Professional Window Tint Removal in Basking Ridge",
        "Is your old window tint bubbling, peeling, or turning purple? Trying to scrape it yourself risks damaging rear defroster lines and glass surfaces. CleanWorx provides professional steam-assisted tint removal starting at $20 per window at our Basking Ridge studio. Replace failing film with premium heat-rejecting Ceramic tint.",
        "Foto del proceso de desprendimiento de film viejo con vapor sin dañar líneas térmicas de la luneta trasera.",
        "Learn more",
        "https://www.cleanworxnj.com/window-tinting?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_window_tint_removal_service_20261228",
        "Verificar que la tarifa base comience en $20 por ventana según el catálogo oficial de servicios.",
        "Polarizado automotriz (Window Tinting)"
    ),
    # Week 12: 2027-01-04
    (
        12, "2027-01-04",
        "Arrancar el año con un plan de mantenimiento y detallado completo",
        "auto detailing service",
        "Start the New Year with a Full Vehicle Detail",
        "Give your daily driver the fresh start it deserves this January. CleanWorx full detailing packages combine comprehensive interior deep cleaning with complete exterior decontamination to eliminate accumulated grit and winter grime. Drop off at our 19 E. Henry Street shop or inquire about our self-contained mobile appointment availability.",
        "Foto panorámica de un vehículo completo recién terminado dentro del taller bajo iluminación profesional.",
        "Book",
        "https://www.cleanworxnj.com/?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_new_year_full_detail_20270104",
        "Probar el enlace al sitio web y confirmar que la información de paquetes esté al día.",
        "Taller, equipo y reputación (Basking Ridge Studio & Trust)"
    ),
    # Week 13: 2027-01-11
    (
        13, "2027-01-11",
        "Ofrecer tratamiento purificador de olores para vehículos cerrados en pleno invierno",
        "smoke smell removal car",
        "Studio-Based Car Odor Treatment & Air Purification",
        "Lingering tobacco smoke, food odors, and stale heater smells trapped in vehicle fabrics won't disappear with a cardboard air freshener. CleanWorx offers dedicated air purification at our Basking Ridge studio. Combined with deep steam cleaning and interior extraction, we tackle odor sources at the fiber level. Studio-only appointment: $125 standalone or $75 add-on.",
        "Foto del interior del vehículo durante el proceso de purificación en el taller cerrado de Basking Ridge.",
        "Learn more",
        "https://www.cleanworxnj.com/car-odor-treatment?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_smoke_odor_treatment_20270111",
        "No hacer promesas de eliminación absoluta de moho ni afirmaciones médicas o de sanitización.",
        "Servicios suplementarios (Odor, Engine, Headlights)"
    ),
    # Week 14: 2027-01-18
    (
        14, "2027-01-18",
        "Explicar el daño de la sal de carretera de NJ y cómo neutralizarla con lavado a mano",
        "exterior detailing",
        "Winter Salt Removal & Exterior Paint Decontamination",
        "Road salt and liquid calcium chloride used across New Jersey highways create caustic chemical reactions that degrade paint clear coats and etch wheels. CleanWorx exterior detailing safely dissolves road grime and salt crust using pH-balanced soaps, wheel well flushes, and ceramic spray sealants. Protect your investment this winter at 19 E. Henry St.",
        "Foto de rines y pasarruedas cubiertos de salmuera siendo enjuagados a presión con agua tibia desionizada.",
        "Book",
        "https://www.cleanworxnj.com/exterior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_winter_salt_decontamination_20270118",
        "Aclarar que en invierno las citas de taller en 19 E. Henry St aseguran un ambiente climatizado.",
        "Detallado exterior (Exterior Detailing)"
    ),
    # Week 15: 2027-01-25
    (
        15, "2027-01-25",
        "Educar sobre las etapas de pulido de pintura (1-step vs multi-stage)",
        "car polishing",
        "Multi-Stage Car Polishing & Clear Coat Refinement",
        "Did improper winter snow brushing leave micro-scratches on your vehicle's paint? Machine polishing restores true depth and gloss by removing microscopic defects. CleanWorx provides single-stage gloss enhancements as well as multi-stage paint correction tailored to your vehicle's specific clear coat hardness. Consult with Vito and the team at our Basking Ridge studio.",
        "Foto de acción de un pulidor trabajando con máquina rotorbital y pad de microfibra sobre un panel oscuro.",
        "Learn more",
        "https://www.cleanworxnj.com/paint-correction?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_car_polishing_stages_20270125",
        "Destacar que el precio final y número de etapas se determinan tras inspección física del vehículo.",
        "Corrección de pintura (Paint Correction)"
    ),
    # Week 16: 2027-02-01
    (
        16, "2027-02-01",
        "Captar dueños de mascotas con necesidad de extracción de pelos de perro/gato incrustados",
        "pet hair removal car",
        "Specialized Pet Hair Removal & Interior Detailing",
        "Loving your pets shouldn't mean sacrificing your car's interior. Pet hair weaves stubbornly into car carpets and seat fabrics where standard vacuums can't reach. CleanWorx uses specialized rubber stone tools, compressed air purging, and high-suction extraction to completely lift embedded pet hair. Bring your vehicle to 19 E. Henry Street in Basking Ridge.",
        "Foto de detalle de un asiento trasero o maletero lleno de pelos de perro antes y completamente impecable después.",
        "Book",
        "https://www.cleanworxnj.com/interior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_pet_hair_removal_service_20270201",
        "Confirmar si aplica tarifa adicional por suciedad extrema o pelo excesivo según catálogo.",
        "Detallado interior (Interior Detailing)"
    ),
    # Week 17: 2027-02-08
    (
        17, "2027-02-08",
        "Clarificar la diferencia entre lámina Carbon y Ceramic y cómo se cotiza",
        "windows tinted price",
        "Carbon vs Ceramic Window Tint: Clear Pricing & Options",
        "Choosing the right automotive window film depends on your priorities. CleanWorx offers Carbon film for sleek factory looks and glare reduction (2-year warranty) and premium Ceramic film for industry-leading infrared heat rejection and UV blocking (10-year warranty). Transparent pricing by vehicle class available at our Basking Ridge shop. Inquire online.",
        "Infografía o foto comparando muestras de film Carbon y Ceramic en el taller con medidor térmico.",
        "Learn more",
        "https://www.cleanworxnj.com/window-tinting?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_window_tint_pricing_guide_20270208",
        "Verificar que los precios publicados coincidan con la tabla oficial en /window-tinting.",
        "Polarizado automotriz (Window Tinting)"
    ),
    # Week 18: 2027-02-15
    (
        18, "2027-02-15",
        "Resolver la duda frecuente sobre cuánto cuesta y por qué vale la pena un coating cerámico",
        "ceramic coating cost",
        "Understanding Ceramic Coating Costs & Long-Term Value",
        "Is ceramic coating worth the investment? Unlike traditional waxes that wash off in weeks, a professionally applied System X coating forms a semi-permanent sacrificial barrier that lasts years. CleanWorx packages include thorough surface decontamination and paint prep to lock in a brilliant showroom shine. Learn about our 1, 3, and 6-year options in Basking Ridge.",
        "Foto de vehículo de alta gama terminado reflejando las luces del taller como un espejo.",
        "Learn more",
        "https://www.cleanworxnj.com/ceramic-coating?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_ceramic_coating_cost_value_20270215",
        "Mantener los paquetes oficiales de 1, 3 y 6 años sin inventar garantías ajenas.",
        "Protección cerámica (Ceramic Coating)"
    ),
    # Week 19: 2027-02-22
    (
        19, "2027-02-22",
        "Mostrar la ventaja de agendar detallado en parques empresariales y oficinas de NJ",
        "mobile detailing service",
        "Mobile Detailing at Your Office or Workplace in NJ",
        "Spend your workday being productive while CleanWorx details your car in your company parking lot. Our self-contained mobile unit operates throughout Somerset, Morris, and Union counties, carrying spot-free water and generator power. Come out to a spotless, refreshed vehicle at the end of the day. Inquire about corporate and multi-vehicle appointments.",
        "Foto de la unidad móvil atendiendo un vehículo ejecutivo en el estacionamiento de un complejo de oficinas.",
        "Book",
        "https://www.cleanworxnj.com/mobile-auto-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_mobile_detailing_workplace_20270222",
        "Recordar tarifa móvil de $50 si el subtotal es menor a $400; sin tarifa para $400+.",
        "Servicio móvil a domicilio (Mobile Detailing)"
    ),
    # Week 20: 2027-03-01
    (
        20, "2027-03-01",
        "Anticipar la temporada de primavera para retirar todo residuo invernal",
        "car detailing services",
        "Plan Your Post-Winter Detail at CleanWorx Basking Ridge",
        "As winter winds down in New Jersey, road salt, calcium chloride, and slush leave vehicles looking tired and gritty. CleanWorx provides comprehensive auto detailing services that restore both interior hygiene and exterior gloss. Secure your pre-spring appointment early at our 19 E. Henry Street shop or reserve our mobile detailing unit.",
        "Collage limpio de dos fotos: limpieza profunda de interiores y acabado brillante de pintura exterior.",
        "Book",
        "https://www.cleanworxnj.com/?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_pre_spring_detailing_services_20270301",
        "Verificar que el calendario de citas de Square refleje los horarios de marzo.",
        "Taller, equipo y reputación (Basking Ridge Studio & Trust)"
    ),
    # Week 21: 2027-03-08
    (
        21, "2027-03-08",
        "Promover la limpieza del vano motor tras el invierno",
        "engine bay cleaning",
        "Professional Engine Bay Cleaning & Degreasing",
        "Road salt, motor oil vapors, and winter dust accumulate in your vehicle's engine compartment over the colder months. CleanWorx uses precision steam cleaning, specialized degreasers, and delicate brush agitation—protecting sensitive electronics and alternator connections. Available for $75 as an add-on or $125 standalone at our Basking Ridge studio.",
        "Foto antes y después del vano motor: mitad izquierda sucia con polvo y sal, mitad derecha vestida y brillante.",
        "Learn more",
        "https://www.cleanworxnj.com/add-ons?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_engine_bay_cleaning_spring_20270308",
        "Confirmar los precios oficiales de catálogo: $75 add-on / $125 standalone.",
        "Servicios suplementarios (Odor, Engine, Headlights)"
    ),
    # Week 22: 2027-03-15
    (
        22, "2027-03-15",
        "Destacar la higiene de vapor y cuidado de cuero en el interior",
        "car steam cleaning",
        "High-Temperature Steam Cleaning & Leather Conditioning",
        "High-temperature steam sanitizes door cards, dashboard vents, center consoles, and cupholders without harsh chemical residue. CleanWorx pairs precision steam detailing with pH-neutral leather cleansers and nourishing conditioners that restore suppleness and prevent premature cracking. Drop off your vehicle at 19 E. Henry Street in Basking Ridge.",
        "Foto de detalle de vapor saliendo de boquilla fina sobre rieles de asiento y rejillas de aire acondicionado.",
        "Book",
        "https://www.cleanworxnj.com/interior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_steam_cleaning_leather_care_20270315",
        "Asegurar que se mencionen limpiadores con pH neutro aptos para cuero automotriz.",
        "Detallado interior (Interior Detailing)"
    ),
    # Week 23: 2027-03-22
    (
        23, "2027-03-22",
        "Anunciar la temporada de corrección de pintura para eliminar rayas del invierno",
        "scratch removal car",
        "Spring Scratch & Swirl Removal at CleanWorx",
        "Strong spring sunlight reveals every scratch, swirl mark, and buffer trail on your vehicle's exterior. CleanWorx specializes in assessment-led paint correction starting at $350. We evaluate clear coat thickness with digital gauges to safely level paint defects and deliver a mirror-like finish. Visit our 19 E. Henry Street shop in Basking Ridge.",
        "Foto del medidor de espesor de pintura apoyado sobre el panel antes de pulir.",
        "Learn more",
        "https://www.cleanworxnj.com/paint-correction?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_spring_scratch_removal_20270322",
        "Reafirmar que la corrección de pintura se realiza en taller tras inspección previa.",
        "Corrección de pintura (Paint Correction)"
    ),
    # Week 24: 2027-03-29
    (
        24, "2027-03-29",
        "Posicionar polarizado cerámico ante la llegada del calor primaveral y radiación solar",
        "where can i get my car windows tinted",
        "Looking for Trusted Window Tinting Near Basking Ridge?",
        "Wondering where to get your car windows tinted by experienced installers? CleanWorx provides professional automotive window tinting at our dedicated shop at 19 E. Henry Street in Basking Ridge, NJ. We use precision plotter-cut patterns for bubble-free, gapless fitment on all vehicle classes. Ask about our 10-year warranty Ceramic film options.",
        "Foto del plotter cortando el film con precisión digital para un modelo específico.",
        "Get directions",
        "https://maps.google.com/?q=CleanWorx+Auto+Detailing+19+E+Henry+Street+Basking+Ridge+NJ+07920",
        "Verificar que la dirección y el enlace en Maps apunten al local de Basking Ridge.",
        "Polarizado automotriz (Window Tinting)"
    ),
    # Week 25: 2027-04-05
    (
        25, "2027-04-05",
        "Promover la protección UV y contra savia/ácidos antes del verano",
        "ceramic paint protection",
        "Protect Your Paint from Spring Tree Sap & Summer UV",
        "Tree sap, bird droppings, and intense UV rays can etch unprotected automotive paint in just a few days. System X ceramic coatings form an inorganic ceramic shield that repels organic contaminants and maintains deep gloss through seasonal extremes. CleanWorx offers certified ceramic applications in Basking Ridge. Book your paint inspection today.",
        "Foto de pintura negra protegida con cerámico con reflejo cristalino de árboles y cielo.",
        "Learn more",
        "https://www.cleanworxnj.com/ceramic-coating?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_ceramic_paint_protection_spring_20270405",
        "Comprobar que se mencione la certificación System X.",
        "Protección cerámica (Ceramic Coating)"
    ),
    # Week 26: 2027-04-12
    (
        26, "2027-04-12",
        "Ofrecer limpieza interior a domicilio durante los días templados de primavera",
        "mobile interior car detailing",
        "Mobile Interior Detailing at Your Home Driveway",
        "Enjoy a fresh, immaculate cabin without leaving your house. CleanWorx dispatches self-contained mobile detailing units across Basking Ridge, Bernardsville, Bedminster, and surrounding Somerset/Morris towns. We extract carpets, shampoo seats, detail plastics, and clean glass right in your driveway. Book an appointment online.",
        "Foto de técnico operando extractor en el interior de una SUV en la entrada de una casa.",
        "Book",
        "https://www.cleanworxnj.com/mobile-auto-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_mobile_interior_detailing_driveway_20270412",
        "Recordar: tarifa de $50 solo si el subtotal de servicios móviles es menor a $400.",
        "Servicio móvil a domicilio (Mobile Detailing)"
    ),
    # Week 27: 2027-04-19
    (
        27, "2027-04-19",
        "Aclarar el valor de un sellador sintético / cera aplicada a mano vs lavado automático",
        "car waxing",
        "Hand Waxing & Long-Lasting Synthetic Sealants",
        "While quick car washes spray temporary silicones that wash away in the rain, professional hand waxing and synthetic polymer sealants create lasting gloss and surface protection. At CleanWorx, every exterior detail includes thorough chemical and mechanical decontamination before wax or sealant application. Book your exterior refresh in Basking Ridge.",
        "Foto de aplicador de cera/sellador siendo retirado con paño de microfibra afelpado revelando brillo húmedo.",
        "Book",
        "https://www.cleanworxnj.com/exterior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_car_waxing_exterior_detail_20270419",
        "No presentar 'car waxing' como servicio aislado de lavado regular; encuadrar como parte de detallado exterior.",
        "Detallado exterior (Exterior Detailing)"
    ),
    # Week 28: 2027-04-26
    (
        28, "2027-04-26",
        "Construir confianza conociendo al equipo: Vito, Melqui y Hemza",
        "professional auto detailing",
        "Meet the CleanWorx Detailing Team in Basking Ridge",
        "At CleanWorx, your vehicle is never handed over to untrained temporary labor. Vito DeGironimo started CleanWorx in 2019, growing it into our dedicated 19 E. Henry Street shop. Today, Vito, Melqui Pichardo, and Hemza Nasser personally execute every correction, coating, and interior detail. Experience craftsmanship where every detail matters.",
        "Foto del equipo (Vito, Melqui, Hemza) frente al taller con uniformes profesionales y autos de clientes.",
        "Learn more",
        "https://www.cleanworxnj.com/about?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_meet_cleanworx_team_20270426",
        "Mantener los nombres reales confirmados por onboarding: Vito DeGironimo, Melqui Pichardo, Hemza Nasser.",
        "Taller, equipo y reputación (Basking Ridge Studio & Trust)"
    ),
    # Week 29: 2027-05-03
    (
        29, "2027-05-03",
        "Responder dudas frecuentes sobre polarizado en New Jersey",
        "window tint in new jersey",
        "Window Tint in New Jersey: Comfort, Privacy & Clarity",
        "Adding window tint to your car improves driving comfort dramatically as temperatures climb in NJ. CleanWorx offers high-clarity Carbon and Ceramic film that blocks solar heat and harmful rays while preserving crystal-clear optical visibility day and night. Visit our Basking Ridge studio to inspect tint darkness levels in person before installation.",
        "Foto de un cliente comparando muestras de tonos de polarizado (5%, 20%, 35%, 50%) en el taller.",
        "Learn more",
        "https://www.cleanworxnj.com/window-tinting?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_window_tint_new_jersey_options_20270503",
        "Evitar dar asesoría legal absoluta o interpretaciones de normativas; enfocarse en especificaciones técnicas de películas.",
        "Polarizado automotriz (Window Tinting)"
    ),
    # Week 30: 2027-05-10
    (
        30, "2027-05-10",
        "Promover remoción de manchas de café, comida y barro de primavera en asientos",
        "car seat stain removal",
        "Professional Car Seat Stain Removal & Shampooing",
        "Spilled drinks, coffee drips, and water marks take away from your car's interior appeal. CleanWorx employs enzyme-based pre-treatments, gentle rotary brush agitation, and hot water extraction to lift deep-seated stains from fabric seats and floor mats. Schedule an interior restoration at 19 E. Henry Street in Basking Ridge.",
        "Foto de antes y después de un asiento de tela beige con mancha de café eliminada.",
        "Book",
        "https://www.cleanworxnj.com/interior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_car_seat_stain_removal_20270510",
        "Comprobar que no se hagan promesas irrealistas en telas descoloridas o quemadas.",
        "Detallado interior (Interior Detailing)"
    ),
    # Week 31: 2027-05-17
    (
        31, "2027-05-17",
        "Mostrar cómo un pulido restaura el brillo de vehículos oscuros",
        "paint polishing",
        "Restoring Depth & Gloss Through Machine Paint Polishing",
        "Does your paint look hazy or gray in direct sunlight? Paint polishing refines the clear coat surface, eliminating micro-marring and oxidation that rob your car of its genuine color depth. CleanWorx pairs micro-finishing polishes with Italian dual-action machines for a flawless, holograph-free result. Available at our Basking Ridge shop starting from $350.",
        "Video corto mostrando reflejos nítidos de iluminación LED tras la pasada final de pulido.",
        "Book",
        "https://www.cleanworxnj.com/paint-correction?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_paint_polishing_refinement_20270517",
        "Mantener la indicación de precio base desde $350 según confirmación del propietario.",
        "Corrección de pintura (Paint Correction)"
    ),
    # Week 32: 2027-05-24
    (
        32, "2027-05-24",
        "Captar limpiezas completas antes del fin de semana largo de Memorial Day",
        "complete car detailing",
        "Memorial Day Road Trip Prep: Complete Car Detailing",
        "Heading down the Shore or hitting the highway for Memorial Day? Travel in comfort with a fully detailed vehicle. Our complete detail packages thoroughly clean, condition, and protect your interior while giving your paint a glossy, contaminant-free finish. Reserve your drop-off time at 19 E. Henry Street or request our mobile detailing unit.",
        "Foto de SUV o sedán familiar impecable listo para viajar estacionado frente al taller.",
        "Book",
        "https://www.cleanworxnj.com/?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_memorial_day_complete_detail_20270524",
        "Comprobar la disponibilidad de citas de cara al fin de semana festivo.",
        "Taller, equipo y reputación (Basking Ridge Studio & Trust)"
    ),
    # Week 33: 2027-05-31
    (
        33, "2027-05-31",
        "Explicar el rechazo de calor por rayos infrarrojos del polarizado cerámico en verano",
        "tint car windows near me",
        "Beat the Summer Heat with Ceramic Window Tinting",
        "Summer temperatures turn vehicle interiors into ovens. CleanWorx Ceramic window film blocks up to 99% of harmful UV rays and drastically rejects infrared heat, keeping your cabin noticeably cooler and easing the strain on your air conditioning. Professionally installed inside our clean Basking Ridge studio with a 10-year warranty.",
        "Foto térmica o demostración de lámpara de calor infrarrojo mostrando el bloqueo térmico del film cerámico.",
        "Learn more",
        "https://www.cleanworxnj.com/window-tinting?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_ceramic_tint_summer_heat_20270531",
        "Recordar la garantía oficial de 10 años en Ceramic film.",
        "Polarizado automotriz (Window Tinting)"
    ),
    # Week 34: 2027-06-07
    (
        34, "2027-06-07",
        "Explicar cómo mantener un auto tratado con cerámico para maximizar su hidrofobia",
        "automotive ceramic coating",
        "How to Maintain Your Automotive Ceramic Coating",
        "A ceramic coating makes washing your vehicle effortless, but proper care preserves its performance. Avoid automatic abrasive car washes and stick to pH-neutral hand washes with soft microfiber mitts. CleanWorx provides System X ceramic installations and maintenance washes for clients across Somerset County. Learn more about protecting your investment.",
        "Video de agua resbalando instantáneamente de un panel con recubrimiento cerámico durante un lavado de mantenimiento.",
        "Learn more",
        "https://www.cleanworxnj.com/ceramic-coating?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_ceramic_coating_maintenance_20270607",
        "Confirmar que las recomendaciones de mantenimiento coincidan con las directrices de System X.",
        "Protección cerámica (Ceramic Coating)"
    ),
    # Week 35: 2027-06-14
    (
        35, "2027-06-14",
        "Destacar la facilidad de reservar servicio móvil durante los fines de semana de verano",
        "mobile car detailing",
        "Professional Detailing Right in Your Driveway This Summer",
        "No need to spend your Saturday waiting in a noisy car wash lobby. CleanWorx brings premium auto detailing right to your home driveway in Basking Ridge, Bridgewater, Warren, and nearby towns. Our mobile rig arrives with pure deionized water and power, ready to leave your car spotless inside and out. Book online.",
        "Foto del técnico secando con aire comprimido un auto en un hermoso driveway residencial de NJ.",
        "Book",
        "https://www.cleanworxnj.com/mobile-auto-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_summer_mobile_detailing_driveway_20270614",
        "Verificar que el cliente entienda la regla de los $50 de cargo móvil si no alcanza los $400.",
        "Servicio móvil a domicilio (Mobile Detailing)"
    ),
    # Week 36: 2027-06-21
    (
        36, "2027-06-21",
        "Promover hidratación y protección de tapicería de cuero contra el sol del verano",
        "auto interior detailing",
        "Protect Leather Seats from Summer Sun & Heat Damage",
        "Baking summer sun dries out leather seats, leading to stiffness, fading, and cracks over time. CleanWorx interior detailing cleans deeply into leather grain pores before infusing breathable UV conditioners that restore natural suppleness. Drop off your car at 19 E. Henry Street in Basking Ridge for expert interior preservation.",
        "Foto en primer plano de un asiento de cuero con acabado mate suave y limpio (sin brillo grasoso barato).",
        "Book",
        "https://www.cleanworxnj.com/interior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_leather_protection_summer_20270621",
        "Resaltar que el acabado final es mate y sedoso, no aceitoso ni resbaladizo.",
        "Detallado interior (Interior Detailing)"
    ),
    # Week 37: 2027-06-28
    (
        37, "2027-06-28",
        "Mostrar proyectos de alta gama realizados en el taller de CleanWorx para generar autoridad",
        "luxury car detailing",
        "Luxury & Exotic Vehicle Detailing in Basking Ridge",
        "From Porsche 911s and BMW M series to everyday luxury SUVs, CleanWorx handles high-value vehicles with meticulous care. We use non-acidic wheel cleaners, soft boars-hair brushes for emblems, and swirl-free wash methods. Explore our completed project gallery online or schedule an appointment at our 19 E. Henry Street facility.",
        "Foto de un auto deportivo o sedán de lujo en el estudio de CleanWorx con iluminación de detalle.",
        "Learn more",
        "https://www.cleanworxnj.com/our-work?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_luxury_car_detailing_portfolio_20270628",
        "Asegurar que las fotos correspondan a proyectos documentados de CleanWorx y se cubran placas si es requerido.",
        "Taller, equipo y reputación (Basking Ridge Studio & Trust)"
    ),
    # Week 38: 2027-07-05
    (
        38, "2027-07-05",
        "Tratar la acumulación de insectos y asfalto derretido típica de los viajes de julio",
        "exterior auto detailing",
        "Safe Bug Splatter & Road Tar Removal for Your Paint",
        "Highway driving in July leaves acidic bug splatters and sticky road tar on your front bumper and rocker panels. Scrubbing them with standard sponges scratches your clear coat. CleanWorx uses specialized citrus bug dissolvers and safe tar removers to melt contaminants without harming paint. Visit 19 E. Henry Street for expert exterior care.",
        "Foto del frontal del vehículo con limpiador de insectos actuando y enjuagándose sin marcas.",
        "Book",
        "https://www.cleanworxnj.com/exterior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_bug_tar_removal_exterior_20270705",
        "Recordar que la remoción de bichos no debe frotarse en seco para no generar micro-rayas.",
        "Detallado exterior (Exterior Detailing)"
    ),
    # Week 39: 2027-07-12
    (
        39, "2027-07-12",
        "Ofrecer tratamiento contra el olor a humedad o moho del aire acondicionado en verano",
        "car odor removal",
        "Fresh Cabin Air: Studio-Based Car Odor Treatment",
        "Does your vehicle's air conditioning blow a damp, musty smell when turned on in humid summer weather? CleanWorx offers specialized air purification at our Basking Ridge studio. Combined with cabin filter inspection and interior steam sanitization, we neutralize trapped moisture smells. Dedicated service available from $125 standalone or $75 as an add-on.",
        "Foto de la cabina limpia con el equipo purificador de aire trabajando en el taller cerrado.",
        "Learn more",
        "https://www.cleanworxnj.com/car-odor-treatment?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_car_ac_odor_treatment_20270712",
        "Evitar promesas de remediación médica de moho; enfocarse en purificación de aire y eliminación de olores.",
        "Servicios suplementarios (Odor, Engine, Headlights)"
    ),
    # Week 40: 2027-07-19
    (
        40, "2027-07-19",
        "Mostrar el efecto hidrofóbico extremo y facilidad de lavado de un auto con coating",
        "ceramic coating car",
        "Experience Extreme Water Beading with System X Ceramic",
        "Watch rain, mud, and road grime slide effortlessly off your paint. System X ceramic coatings create an ultra-low surface tension that repels liquids and prevents mineral water spots from etching your finish. Drop off your vehicle at CleanWorx in Basking Ridge for multi-year paint protection installed by certified technicians.",
        "Video en cámara lenta o foto de gotas de agua saliendo disparadas del capó cerámico con un soplador de aire.",
        "Learn more",
        "https://www.cleanworxnj.com/ceramic-coating?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_ceramic_coating_water_beading_20270719",
        "Confirmar que el video o foto muestre trabajo real del taller.",
        "Protección cerámica (Ceramic Coating)"
    ),
    # Week 41: 2027-07-26
    (
        41, "2027-07-26",
        "Promover el cambio de polarizado defectuoso por láminas cerámicas modernas",
        "window tint removal cost",
        "Replace Damaged, Peeling Tint in Basking Ridge",
        "Distorted or bubbling window film looks unsightly and severely impairs your driving visibility. CleanWorx offers careful steam-assisted tint removal starting from $20 per window at our 19 E. Henry Street shop. Upgrade to our optical-clarity Ceramic film to restore both looks and superior thermal performance.",
        "Foto de antes (ventana con film lleno de burbujas moradas) y después (vidrio limpio con film cerámico nítido).",
        "Learn more",
        "https://www.cleanworxnj.com/window-tinting?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_tint_removal_replacement_cost_20270726",
        "Verificar que el cliente conozca la tarifa inicial de $20 por ventana de remoción.",
        "Polarizado automotriz (Window Tinting)"
    ),
    # Week 42: 2027-08-02
    (
        42, "2027-08-02",
        "Limpiar arena, manchas de protector solar y humedad de vacaciones de playa",
        "deep clean car interior",
        "Post-Beach Interior Deep Clean & Sand Extraction",
        "Summer beach trips leave fine sand lodged in carpet weave and white sunscreen smudges across door armrests and leather seats. CleanWorx interior detailing uses high-pressure air agitation, commercial vacuuming, and specialized upholstery shampoos to extract stubborn sand and restore cabin freshness. Book at 19 E. Henry Street in Basking Ridge.",
        "Foto de arena siendo extraída de los rieles de los asientos y piso con herramienta de aire comprimido (Tornador).",
        "Book",
        "https://www.cleanworxnj.com/interior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_post_beach_interior_deep_clean_20270802",
        "Probar que el botón de reserva dirija a la página interna correspondiente.",
        "Detallado interior (Interior Detailing)"
    ),
    # Week 43: 2027-08-09
    (
        43, "2027-08-09",
        "Promover detallado de múltiples autos en la misma residencia o vecindario",
        "mobile detailing",
        "Detail Multiple Family Vehicles in One Mobile Visit",
        "Have two or three vehicles at home that need a professional touch? Schedule CleanWorx for a combined mobile detailing appointment. Our self-contained van comes fully equipped to handle sedans, SUVs, and trucks right in your driveway. Reach the $400 subtotal threshold and enjoy zero mobile service fees! Serving Basking Ridge and nearby communities.",
        "Foto de dos vehículos familiares (ej. una SUV y un sedán) estacionados juntos siendo detallados por la unidad móvil.",
        "Book",
        "https://www.cleanworxnj.com/mobile-auto-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_multi_vehicle_mobile_detailing_20270809",
        "Destacar la regla de ahorro: subtotal de $400+ elimina la tarifa móvil de $50.",
        "Servicio móvil a domicilio (Mobile Detailing)"
    ),
    # Week 44: 2027-08-16
    (
        44, "2027-08-16",
        "Mostrar la transformación de color y brillo profundo antes de terminar el verano",
        "paint restoration car",
        "Revive Dull & Oxidized Paint at CleanWorx",
        "Has your vehicle lost its glossy reflection under the summer sun? Paint correction gently levels the microscopic valleys and peaks in your clear coat, revealing rich, deep color and sharp reflections. CleanWorx offers tailored correction packages starting from $350 at our Basking Ridge studio. Schedule your paint assessment with our certified team.",
        "Foto antes y después de panel lateral mostrando reflejo borroso antes y espejo nítido después.",
        "Learn more",
        "https://www.cleanworxnj.com/paint-correction?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_paint_restoration_gloss_revival_20270816",
        "Aclarar que rayones profundos que exponen la imprimación requieren pintura de retoque, no pulido.",
        "Corrección de pintura (Paint Correction)"
    ),
    # Week 45: 2027-08-23
    (
        45, "2027-08-23",
        "Limpiar vehículos familiares antes del inicio del año escolar y las rutinas diarias",
        "car detailing near me",
        "Back-to-School Vehicle Detailing in Basking Ridge",
        "Summer road trips and beach days are winding down—get your family vehicle refreshed and organized for school drop-offs and daily commutes. CleanWorx full detailing removes accumulated crumbs, spills, and exterior road dust. Drop off your car at 19 E. Henry Street in Basking Ridge or reserve a convenient mobile appointment.",
        "Foto de una minivan o SUV familiar con interior perfectamente ordenado y brillante.",
        "Book",
        "https://www.cleanworxnj.com/?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_back_to_school_car_detailing_20270823",
        "Revisar disponibilidad de turnos en las semanas previas a septiembre.",
        "Taller, equipo y reputación (Basking Ridge Studio & Trust)"
    ),
    # Week 46: 2027-08-30
    (
        46, "2027-08-30",
        "Educar al cliente sobre por qué la cera se derrite en verano y el cerámico resiste altas temperaturas",
        "ceramic coating service",
        "Why Ceramic Coating Outlasts Traditional Wax",
        "Natural carnauba waxes break down and melt at temperatures above 140°F—easily reached on dark car hoods in August. System X ceramic coatings withstand extreme heat up to 1000°F without losing their hydrophobic bond or gloss. Protect your car for years, not weeks, with professional ceramic coating at CleanWorx in Basking Ridge, NJ.",
        "Foto de panel caliente bajo el sol con el agua escurriendo limpiamente.",
        "Learn more",
        "https://www.cleanworxnj.com/ceramic-coating?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_ceramic_coating_vs_wax_durability_20270830",
        "Asegurar que los datos técnicos de temperatura sean precisos y representativos de System X.",
        "Protección cerámica (Ceramic Coating)"
    ),
    # Week 47: 2027-09-06
    (
        47, "2027-09-06",
        "Aprovechar las semanas de septiembre para polarizar sin colas de verano",
        "window tint service near me",
        "Autumn Window Tinting at CleanWorx Studio",
        "September brings clear skies and lower sun angles that create sharp glare on your morning and evening drive. CleanWorx offers high-performance Carbon and Ceramic window film installed inside our climate-controlled Basking Ridge shop. Enjoy enhanced driving comfort, privacy, and long-term interior UV protection. Call or book online.",
        "Foto de vehículo con lunas tintadas estacionado con follaje otoñal temprano de fondo en NJ.",
        "Get directions",
        "https://maps.google.com/?q=CleanWorx+Auto+Detailing+19+E+Henry+Street+Basking+Ridge+NJ+07920",
        "Comprobar el funcionamiento del botón 'Get directions' hacia la dirección de Basking Ridge.",
        "Polarizado automotriz (Window Tinting)"
    ),
    # Week 48: 2027-09-13
    (
        48, "2027-09-13",
        "Recordar la restauración de faros de cara a las noches más largas de otoño",
        "car headlight restoration",
        "Restore Dull Headlights for Brighter Autumn Driving",
        "Don't wait until cloudy autumn rains and early darkness make driving hazardous. CleanWorx wet sands away yellowed oxidation and seals your polycarbonate lenses with durable UV protection. Available as a convenient $75 add-on to any detailing package or $125 standalone at our 19 E. Henry Street shop in Basking Ridge.",
        "Foto comparativa dividida mostrando el patrón de luz del faro opaco vs faro restaurado iluminando la pared.",
        "Learn more",
        "https://www.cleanworxnj.com/add-ons?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_headlight_restoration_fall_check_20270913",
        "Verificar que los precios reflejen las tarifas de catálogo ($75 add-on, $125 standalone).",
        "Servicios suplementarios (Odor, Engine, Headlights)"
    ),
    # Week 49: 2027-09-20
    (
        49, "2027-09-20",
        "Preparar el interior del auto para el cambio de temporada",
        "car shampoo service",
        "Full Interior Shampoo & Fabric Reconditioning",
        "Remove summer sweat, beach residue, and spilled drinks before cold weather sets in. CleanWorx provides thorough interior shampooing with heated extraction, leaving fabrics sanitized, soft, and fast-drying. Drop off your vehicle at 19 E. Henry Street in Basking Ridge or inquire about our self-contained mobile detailing unit.",
        "Foto de técnico pasando máquina de inyección-extracción sobre tapicería de asientos.",
        "Book",
        "https://www.cleanworxnj.com/interior-detailing?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_interior_shampoo_fabric_clean_20270920",
        "Recordar ventilación y secado completo en cabina antes de la entrega al cliente.",
        "Detallado interior (Interior Detailing)"
    ),
    # Week 50: 2027-09-27
    (
        50, "2027-09-27",
        "Corregir daños leves acumulados en verano antes del mal tiempo invernal",
        "auto paint correction",
        "Professional Auto Paint Correction in Basking Ridge",
        "Summer road trips and rapid car washes often leave faint scratches and swirl marks across your vehicle's clear coat. CleanWorx machine paint correction safely removes imperfections and restores optical clarity before the winter elements hit New Jersey. Consult with Vito and the team at 19 E. Henry Street to evaluate your vehicle's paint.",
        "Foto de inspección de pintura con lámpara LED mostrando el acabado libre de hologramas.",
        "Learn more",
        "https://www.cleanworxnj.com/paint-correction?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_auto_paint_correction_basking_ridge_20270927",
        "Destacar que el alcance y precio exacto se definen con base en las necesidades del cliente y el estado del auto.",
        "Corrección de pintura (Paint Correction)"
    ),
    # Week 51: 2027-10-04
    (
        51, "2027-10-04",
        "Reforzar la presencia local en Basking Ridge y condados circundantes",
        "auto detailing nj",
        "Local Auto Detailing Excellence in Somerset County, NJ",
        "Proudly serving Basking Ridge, Bernardsville, Bedminster, Warren, and Bridgewater. CleanWorx combines studio craftsmanship at 19 E. Henry Street with the convenience of our fully equipped mobile detailing unit. Whether you need ceramic coating, paint correction, window tinting, or deep interior care, we deliver exceptional results. Book online today.",
        "Foto del taller con la bandera local o vista de la calle Henry St con vehículos de clientes atendidos.",
        "Get directions",
        "https://maps.google.com/?q=CleanWorx+Auto+Detailing+19+E+Henry+Street+Basking+Ridge+NJ+07920",
        "Comprobar que los datos NAP coincidan al 100% con la ficha de Google.",
        "Taller, equipo y reputación (Basking Ridge Studio & Trust)"
    ),
    # Week 52: 2027-10-11
    (
        52, "2027-10-11",
        "Cerrar el ciclo anual con un paquete de detallado completo y preparar el vehículo para el invierno",
        "full car detail",
        "Annual Vehicle Reset: Complete Full Car Detail",
        "Complete the full seasonal care cycle for your car. A CleanWorx full detail thoroughly restores every surface: deep interior shampoo, leather conditioning, exterior chemical decontamination, wheel cleaning, and long-lasting paint sealant. Experience why drivers across New Jersey trust CleanWorx. Schedule your appointment at 19 E. Henry Street.",
        "Foto final de entrega de un vehículo completamente detallado con el cliente recibiendo las llaves con satisfacción.",
        "Book",
        "https://www.cleanworxnj.com/?utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content=update_annual_full_car_detail_reset_20271011",
        "Verificar que la cotización de full detail incluya tanto el interior como el exterior según el catálogo.",
        "Taller, equipo y reputación (Basking Ridge Studio & Trust)"
    )
]

print(f"Total posts drafted: {len(posts_data)}")

# Verify all keywords exist in keywords_dict
missing_kws = [p[3] for p in posts_data if p[3] not in keywords_dict]
if missing_kws:
    print("WARNING: Missing keywords in CSV:", missing_kws)
else:
    print("ALL 52 keywords match CSV exactly!")

# --- 3. CREATE WORKBOOK ---
wb = openpyxl.Workbook()

# Sheet 1: Calendario 52 semanas
ws_cal = wb.active
ws_cal.title = "Calendario 52 semanas"

# Sheet 2: Resumen
ws_res = wb.create_sheet(title="Resumen")

# Sheet 3: Keywords Ubersuggest (same name as template)
ws_kw = wb.create_sheet(title="Keywords Ubersuggest")

# --- STYLES & PALETTES (Clean, Luxury Automotive Theme) ---
font_name = "Carlito"
navy_dark = "FF0F172A"       # Slate 900 for main banner
blue_primary = "FF1E40AF"    # Royal Automotive Blue for table headers
blue_soft = "FFEFF6FF"       # Very soft blue for asset link row
zebra_bg = "FFF8FAFC"        # Slate 50 for alternating rows
white_bg = "FFFFFFFF"
gold_accent = "FFB58B2A"     # Subtle gold for special highlights
border_color = "FFE2E8F0"    # Light slate border

thin_border = Border(
    left=Side(style='thin', color=border_color),
    right=Side(style='thin', color=border_color),
    top=Side(style='thin', color=border_color),
    bottom=Side(style='thin', color=border_color)
)

header_border = Border(
    left=Side(style='thin', color="FF1E3A8A"),
    right=Side(style='thin', color="FF1E3A8A"),
    top=Side(style='thin', color="FF1E3A8A"),
    bottom=Side(style='medium', color="FF0F172A")
)

# --- POPULATE SHEET 1: Calendario 52 semanas ---
ws_cal.views.sheetView[0].showGridLines = True
ws_cal.row_dimensions[1].height = 14
ws_cal.row_dimensions[2].height = 36
ws_cal.row_dimensions[3].height = 22
ws_cal.row_dimensions[4].height = 30

# Row 2: Banner Title
ws_cal.merge_cells("A2:K2")
cell_title = ws_cal.cell(2, 1, "Calendario anual de publicaciones para Google Business Profile — CleanWorx Auto Detailing")
cell_title.font = Font(name=font_name, size=15, bold=True, color="FFFFFFFF")
cell_title.fill = PatternFill(fill_type="solid", fgColor=navy_dark)
cell_title.alignment = Alignment(horizontal="center", vertical="center")

# Row 3: Drive Link
ws_cal.merge_cells("A3:K3")
cell_drive = ws_cal.cell(3, 1, "Carpeta de recursos visuales (Drive): https://drive.google.com/drive/folders/cleanworx-gbp-assets")
cell_drive.font = Font(name=font_name, size=10, italic=True, color="FF1E3A8A")
cell_drive.fill = PatternFill(fill_type="solid", fgColor=blue_soft)
cell_drive.alignment = Alignment(horizontal="center", vertical="center")

# Row 4: Column Headers
cal_headers = [
    "Fecha", "Objetivo", "Keyword original", "Volumen",
    "Título GBP", "Texto listo para publicar (inglés)",
    "Idea de foto o video", "Botón", "Destino",
    "Revisión antes de publicar", "Estado"
]

for col_idx, h_text in enumerate(cal_headers, start=1):
    c = ws_cal.cell(4, col_idx, h_text)
    c.font = Font(name=font_name, size=11, bold=True, color="FFFFFFFF")
    c.fill = PatternFill(fill_type="solid", fgColor=blue_primary)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = header_border

# Rows 5 to 56: Posts Data
for idx, p in enumerate(posts_data, start=5):
    # p: (week, date, obj, kw, title, text, photo, btn, dest, rev, cat)
    dt_val = datetime.datetime.strptime(p[1], "%Y-%m-%d")
    vol_val = keywords_dict.get(p[3], 0.0)
    bg_color = zebra_bg if idx % 2 == 0 else white_bg
    fill_row = PatternFill(fill_type="solid", fgColor=bg_color)

    ws_cal.row_dimensions[idx].height = 68

    # Col A: Fecha
    cA = ws_cal.cell(idx, 1, dt_val)
    cA.number_format = 'yyyy-mm-dd'
    cA.alignment = Alignment(horizontal="center", vertical="top")

    # Col B: Objetivo
    cB = ws_cal.cell(idx, 2, p[2])
    cB.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Col C: Keyword original
    cC = ws_cal.cell(idx, 3, p[3])
    cC.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Col D: Volumen
    cD = ws_cal.cell(idx, 4, vol_val)
    cD.number_format = '#,##0'
    cD.alignment = Alignment(horizontal="right", vertical="top")

    # Col E: Título GBP
    cE = ws_cal.cell(idx, 5, p[4])
    cE.font = Font(name=font_name, size=10, bold=True, color="FF0F172A")
    cE.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Col F: Texto listo para publicar
    cF = ws_cal.cell(idx, 6, p[5])
    cF.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Col G: Idea de foto o video
    cG = ws_cal.cell(idx, 7, p[6])
    cG.font = Font(name=font_name, size=9.5, italic=True, color="FF334155")
    cG.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Col H: Botón
    cH = ws_cal.cell(idx, 8, p[7])
    cH.font = Font(name=font_name, size=10, bold=True, color="FF1E40AF")
    cH.alignment = Alignment(horizontal="center", vertical="top")

    # Col I: Destino
    cI = ws_cal.cell(idx, 9, p[8])
    cI.font = Font(name=font_name, size=9, color="FF2563EB", underline="single")
    cI.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Col J: Revisión antes de publicar
    cJ = ws_cal.cell(idx, 10, p[9])
    cJ.font = Font(name=font_name, size=9.5, color="FF475569")
    cJ.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Col K: Estado
    cK = ws_cal.cell(idx, 11, "Borrador")
    cK.font = Font(name=font_name, size=10, bold=True, color="FFD97706") # Amber for draft
    cK.alignment = Alignment(horizontal="center", vertical="top")

    for col_c in range(1, 12):
        c_cell = ws_cal.cell(idx, col_c)
        if not c_cell.font or c_cell.font.name != font_name:
            c_cell.font = Font(name=font_name, size=10)
        c_cell.fill = fill_row
        c_cell.border = thin_border

# Data Validation for Col K (Estado)
dv = DataValidation(type="list", formula1='"Borrador,Listo,Programado,Publicado"', allow_blank=True)
ws_cal.add_data_validation(dv)
dv.add("K5:K56")

# Column widths for Sheet 1
col_widths = {
    'A': 13,
    'B': 28,
    'C': 26,
    'D': 12,
    'E': 34,
    'F': 52,
    'G': 36,
    'H': 15,
    'I': 45,
    'J': 34,
    'K': 14
}
for col_l, w in col_widths.items():
    ws_cal.column_dimensions[col_l].width = w


# --- POPULATE SHEET 2: Resumen ---
ws_res.views.sheetView[0].showGridLines = True
ws_res.column_dimensions['A'].width = 30
ws_res.column_dimensions['B'].width = 22
ws_res.column_dimensions['C'].width = 5
ws_res.column_dimensions['D'].width = 8
ws_res.column_dimensions['E'].width = 52
ws_res.column_dimensions['F'].width = 8
ws_res.column_dimensions['G'].width = 8
ws_res.column_dimensions['H'].width = 8

ws_res.row_dimensions[1].height = 36
ws_res.row_dimensions[3].height = 26
ws_res.row_dimensions[10].height = 26
ws_res.row_dimensions[11].height = 24

# Banner
ws_res.merge_cells("A1:H1")
c_res_title = ws_res.cell(1, 1, "Resumen del calendario GBP — CleanWorx Auto Detailing")
c_res_title.font = Font(name=font_name, size=15, bold=True, color="FFFFFFFF")
c_res_title.fill = PatternFill(fill_type="solid", fgColor=navy_dark)
c_res_title.alignment = Alignment(horizontal="center", vertical="center")

# Table 1: Indicadores (A3:B8)
ws_res.cell(3, 1, "Indicador").font = Font(name=font_name, size=11, bold=True, color="FFFFFFFF")
ws_res.cell(3, 1).fill = PatternFill(fill_type="solid", fgColor=blue_primary)
ws_res.cell(3, 1).alignment = Alignment(horizontal="center", vertical="center")
ws_res.cell(3, 1).border = header_border

ws_res.cell(3, 2, "Valor").font = Font(name=font_name, size=11, bold=True, color="FFFFFFFF")
ws_res.cell(3, 2).fill = PatternFill(fill_type="solid", fgColor=blue_primary)
ws_res.cell(3, 2).alignment = Alignment(horizontal="center", vertical="center")
ws_res.cell(3, 2).border = header_border

indicadores = [
    ("Total de publicaciones", 52),
    ("Frecuencia", "1 por semana"),
    ("Primera publicación", datetime.datetime(2026, 10, 19)),
    ("Última publicación", datetime.datetime(2027, 10, 11)),
    ("Estado inicial", "Borrador")
]

for i, (ind, val) in enumerate(indicadores, start=4):
    c1 = ws_res.cell(i, 1, ind)
    c1.font = Font(name=font_name, size=10, bold=(i==4))
    c1.border = thin_border
    c2 = ws_res.cell(i, 2, val)
    c2.font = Font(name=font_name, size=10, bold=(i==4))
    c2.border = thin_border
    if isinstance(val, datetime.datetime):
        c2.number_format = 'yyyy-mm-dd'
        c2.alignment = Alignment(horizontal="center")
    elif isinstance(val, (int, float)):
        c2.alignment = Alignment(horizontal="center")
    else:
        c2.alignment = Alignment(horizontal="center")

# Table 2: Estados (D3:E7)
ws_res.cell(3, 4, "Estado").font = Font(name=font_name, size=11, bold=True, color="FFFFFFFF")
ws_res.cell(3, 4).fill = PatternFill(fill_type="solid", fgColor=blue_primary)
ws_res.cell(3, 4).alignment = Alignment(horizontal="center", vertical="center")
ws_res.cell(3, 4).border = header_border

ws_res.cell(3, 5, "Cantidad").font = Font(name=font_name, size=11, bold=True, color="FFFFFFFF")
ws_res.cell(3, 5).fill = PatternFill(fill_type="solid", fgColor=blue_primary)
ws_res.cell(3, 5).alignment = Alignment(horizontal="center", vertical="center")
ws_res.cell(3, 5).border = header_border

estados = ["Borrador", "Listo", "Programado", "Publicado"]
for j, est in enumerate(estados, start=4):
    c_e1 = ws_res.cell(j, 4, est)
    c_e1.font = Font(name=font_name, size=10)
    c_e1.border = thin_border
    c_e2 = ws_res.cell(j, 5, f"=COUNTIF('Calendario 52 semanas'!$K$5:$K$56,D{j})")
    c_e2.font = Font(name=font_name, size=10, bold=True)
    c_e2.alignment = Alignment(horizontal="center")
    c_e2.border = thin_border

# Table 3: Categorías (A11:B20)
ws_res.cell(11, 1, "Categoría").font = Font(name=font_name, size=11, bold=True, color="FFFFFFFF")
ws_res.cell(11, 1).fill = PatternFill(fill_type="solid", fgColor=blue_primary)
ws_res.cell(11, 1).alignment = Alignment(horizontal="center", vertical="center")
ws_res.cell(11, 1).border = header_border

ws_res.cell(11, 2, "Publicaciones").font = Font(name=font_name, size=11, bold=True, color="FFFFFFFF")
ws_res.cell(11, 2).fill = PatternFill(fill_type="solid", fgColor=blue_primary)
ws_res.cell(11, 2).alignment = Alignment(horizontal="center", vertical="center")
ws_res.cell(11, 2).border = header_border

categorias_summary = [
    ("Protección cerámica (Ceramic Coating)", 7),
    ("Corrección de pintura (Paint Correction)", 6),
    ("Polarizado automotriz (Window Tinting)", 8),
    ("Detallado interior (Interior Detailing)", 8),
    ("Detallado exterior (Exterior Detailing)", 4),
    ("Servicio móvil a domicilio (Mobile Detailing)", 5),
    ("Servicios suplementarios (Odor, Engine, Headlights)", 5),
    ("Taller, equipo y reputación (Basking Ridge Studio & Trust)", 9),
]

for k, (cat_name, cat_count) in enumerate(categorias_summary, start=12):
    c_c1 = ws_res.cell(k, 1, cat_name)
    c_c1.font = Font(name=font_name, size=10)
    c_c1.border = thin_border
    c_c2 = ws_res.cell(k, 2, cat_count)
    c_c2.font = Font(name=font_name, size=10, bold=True)
    c_c2.alignment = Alignment(horizontal="center")
    c_c2.border = thin_border

# Total row for Categorías
c_tot1 = ws_res.cell(20, 1, "Total")
c_tot1.font = Font(name=font_name, size=10, bold=True)
c_tot1.border = thin_border
c_tot2 = ws_res.cell(20, 2, "=SUM(B12:B19)")
c_tot2.font = Font(name=font_name, size=10, bold=True)
c_tot2.alignment = Alignment(horizontal="center")
c_tot2.border = thin_border

# Table 4: Reglas Operativas (D10:E17)
ws_res.merge_cells("D10:H10")
c_reglas_title = ws_res.cell(10, 4, "Reglas operativas para cada publicación — CleanWorx")
c_reglas_title.font = Font(name=font_name, size=11, bold=True, color="FFFFFFFF")
c_reglas_title.fill = PatternFill(fill_type="solid", fgColor=navy_dark)
c_reglas_title.alignment = Alignment(horizontal="center", vertical="center")

reglas = [
    ("1", "Usar una sola keyword principal por publicación sin forzar texto repetitivo."),
    ("2", "Tratar 'near me' como señal de proximidad: redactar de forma natural integrando Basking Ridge, NJ y condados vecinos (Somerset, Morris, Union)."),
    ("3", "Cumplir estrictamente la taxonomía UTM: utm_source=google&utm_medium=organic&utm_campaign=gbp_posts&utm_content={tipo}_{tema}_{aaaammdd}."),
    ("4", "Priorizar fotos y videos reales de vehículos atendidos en el taller de 19 E. Henry St o en la unidad móvil (sin fotos de stock)."),
    ("5", "Para servicio móvil, aclarar siempre la tarifa de $50 si el subtotal es menor a $400; sin tarifa adicional si es de $400 o más."),
    ("6", "No prometer eliminación del 100% de rayones profundos sin evaluar pintura previa, ni hacer reclamos médicos o absolutos en purificación de aire / olores."),
    ("7", "Enlazar siempre a la página canónica del servicio correspondiente en el sitio web; nunca enlazar directamente a URLs externas no verificadas ni a Square sin medición.")
]

for r_idx, (num, regla_text) in enumerate(reglas, start=11):
    c_r1 = ws_res.cell(r_idx, 4, num)
    c_r1.font = Font(name=font_name, size=10, bold=True)
    c_r1.alignment = Alignment(horizontal="center", vertical="top")
    c_r1.border = thin_border

    # Merge across D:H for the text or keep in E
    ws_res.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=8)
    c_r2 = ws_res.cell(r_idx, 5, regla_text)
    c_r2.font = Font(name=font_name, size=10)
    c_r2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for col_m in range(5, 9):
        ws_res.cell(r_idx, col_m).border = thin_border


# --- POPULATE SHEET 3: Keywords Ubersuggest ---
ws_kw.views.sheetView[0].showGridLines = True
ws_kw.row_dimensions[1].height = 28

kw_headers = [
    "No", "Palabra clave", "Volumen", "CPC",
    "Dificultad pagada", "Dificultad SEO", "Cluster", "Usos en calendario"
]

for col_idx, h_text in enumerate(kw_headers, start=1):
    c = ws_kw.cell(1, col_idx, h_text)
    c.font = Font(name=font_name, size=11, bold=True, color="FFFFFFFF")
    c.fill = PatternFill(fill_type="solid", fgColor=blue_primary)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = header_border

for idx, (no, kw, vol, cpc, pd, sd, cluster) in enumerate(keywords_list, start=2):
    ws_kw.row_dimensions[idx].height = 20
    bg_color = zebra_bg if idx % 2 == 0 else white_bg
    fill_row = PatternFill(fill_type="solid", fgColor=bg_color)

    c1 = ws_kw.cell(idx, 1, no)
    c1.alignment = Alignment(horizontal="center")
    c2 = ws_kw.cell(idx, 2, kw)
    c2.font = Font(name=font_name, size=10, bold=True)
    c2.alignment = Alignment(horizontal="left")
    c3 = ws_kw.cell(idx, 3, vol)
    c3.number_format = '#,##0'
    c3.alignment = Alignment(horizontal="right")
    c4 = ws_kw.cell(idx, 4, cpc)
    c4.number_format = '$#,##0.00'
    c4.alignment = Alignment(horizontal="right")
    c5 = ws_kw.cell(idx, 5, pd)
    c5.alignment = Alignment(horizontal="center")
    c6 = ws_kw.cell(idx, 6, sd)
    c6.alignment = Alignment(horizontal="center")
    c7 = ws_kw.cell(idx, 7, cluster)
    c7.alignment = Alignment(horizontal="left")
    c8 = ws_kw.cell(idx, 8, f"=COUNTIF('Calendario 52 semanas'!$C$5:$C$56,B{idx})")
    c8.font = Font(name=font_name, size=10, bold=True)
    c8.alignment = Alignment(horizontal="center")

    for col_c in range(1, 9):
        cell_c = ws_kw.cell(idx, col_c)
        if not cell_c.font or cell_c.font.name != font_name:
            cell_c.font = Font(name=font_name, size=10)
        cell_c.fill = fill_row
        cell_c.border = thin_border

# Footnote in Keywords sheet
last_row = len(keywords_list) + 3
ws_kw.merge_cells(f"A{last_row}:H{last_row}")
c_foot = ws_kw.cell(last_row, 1, "Fuente: keyword research CLEANWORX v2.csv (141 palabras clave optimizadas para CleanWorx Auto Detailing & Ceramic Coating)")
c_foot.font = Font(name=font_name, size=10, italic=True, color="FF64748B")

kw_col_widths = {
    'A': 8,
    'B': 36,
    'C': 14,
    'D': 12,
    'E': 18,
    'F': 16,
    'G': 28,
    'H': 18
}
for col_l, w in kw_col_widths.items():
    ws_kw.column_dimensions[col_l].width = w

# --- SAVE OUTPUT FILES ---
dest_project = r"c:\Users\Albin Rodriguez\Documents\CleanWorx\docs\seo\Calendario-GBP-CleanWorx-52-Semanas.xlsx"
dest_downloads = r"C:\Users\Albin Rodriguez\Downloads\Calendario-GBP-CleanWorx-52-Semanas.xlsx"

os.makedirs(os.path.dirname(dest_project), exist_ok=True)
wb.save(dest_project)
print(f"Saved project file to: {dest_project}")

try:
    shutil.copyfile(dest_project, dest_downloads)
    print(f"Saved user Downloads copy to: {dest_downloads}")
except Exception as e:
    print(f"Could not copy to Downloads: {e}")

print("GENERATION COMPLETE AND SUCCESSFUL!")
